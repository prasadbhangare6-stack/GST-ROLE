from datetime import datetime, timezone
from pathlib import Path
import json
import shutil

from fastapi import (
    FastAPI,
    UploadFile,
    File,
    Form,
    HTTPException,
    Response,
    Depends,
    Cookie,
)

from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import Workbook
from openpyxl.styles import Font

from .db import init_db, connect
from .ocr import extract_text
from .parser import parse_invoice, GSTIN_RE
from .login import authenticate_user
from .security import hash_password
from .session import create_session
from .dependencies import get_current_user
from .verification import verify_invoice
from pydantic import BaseModel


# ============================================================
# PATHS
# ============================================================

ROOT = Path(__file__).resolve().parent.parent

UPLOADS = ROOT / "data" / "uploads"
EXPORTS = ROOT / "data" / "exports"

UPLOADS.mkdir(parents=True, exist_ok=True)
EXPORTS.mkdir(parents=True, exist_ok=True)


# ============================================================
# APP
# ============================================================

app = FastAPI(title="GSTFlow Local MVP")

init_db()

app.mount(
    "/static",
    StaticFiles(directory=ROOT / "app" / "static"),
    name="static",
)


# ============================================================
# HOME
# ============================================================

@app.get("/")
def home():
    return FileResponse(
        ROOT / "app" / "static" / "index.html"
    )


# ============================================================
# AUTHENTICATION
# ============================================================

@app.post("/api/login")
def login(
    response: Response,
    username: str = Form(...),
    password: str = Form(...),
):
    username = username.strip()

    with connect() as c:
        user, session_token = authenticate_user(
            c,
            username,
            password,
        )

    if not user or not session_token:
        raise HTTPException(
            status_code=401,
            detail="Invalid username or password",
        )

    response.set_cookie(
        key="session",
        value=session_token,
        httponly=True,
        samesite="lax",
        secure=False,
        max_age=12 * 60 * 60,
    )

    return {
        "ok": True,
        "username": user["username"],
        "role": user["role"],
        "organization_id": user["organization_id"],
    }


# ============================================================
# LOGOUT
# ============================================================

@app.post("/api/logout")
def logout(
    response: Response,
    session: str | None = Cookie(default=None),
):
    if session:
        with connect() as c:
            from .session import revoke_session
            revoke_session(c, session)

    response.delete_cookie(
        key="session",
        httponly=True,
        samesite="lax",
        secure=False,
    )

    return {"ok": True}


# ============================================================
# CURRENT USER
# ============================================================

@app.get("/api/me")
def me(current_user=Depends(get_current_user)):
    return {
        "ok": True,
        "user_id": current_user["user_id"],
        "username": current_user["username"],
        "role": current_user["role"],
        "organization_id": current_user["organization_id"],
    }


# ============================================================
# BUSINESSES
# ============================================================

@app.get("/api/businesses")
def businesses(current_user=Depends(get_current_user)):

    with connect() as c:

        # OWNER: see every business inside their organization.
        if current_user["role"] == "OWNER":

            rows = c.execute(
                """
                SELECT *
                FROM businesses
                WHERE organization_id = ?
                ORDER BY name
                """,
                (current_user["organization_id"],)
            ).fetchall()

        # STAFF: see only businesses explicitly assigned to them.
        else:

            rows = c.execute(
                """
                SELECT b.*
                FROM businesses b
                JOIN user_business_access uba
                    ON uba.business_id = b.id
                WHERE b.organization_id = ?
                  AND uba.user_id = ?
                ORDER BY b.name
                """,
                (
                    current_user["organization_id"],
                    current_user["user_id"],
                )
            ).fetchall()

        return [dict(row) for row in rows]


@app.post("/api/businesses")
def create_business(
    name: str = Form(...),
    gstin: str = Form(...),
    current_user=Depends(get_current_user),
):

    if current_user["role"] != "OWNER":
        raise HTTPException(
            status_code=403,
            detail="Only organization owners can create businesses",
        )

    name = name.strip()
    gstin = gstin.strip().upper()

    if not name:
        raise HTTPException(
            status_code=400,
            detail="Business name is required",
        )

    if not gstin:
        raise HTTPException(
            status_code=400,
            detail="GSTIN is required",
        )

    try:

        with connect() as c:

            cur = c.execute(
                """
                INSERT INTO businesses(
                    organization_id,
                    name,
                    gstin
                )
                VALUES (?, ?, ?)
                """,
                (
                    current_user["organization_id"],
                    name,
                    gstin,
                )
            )

            c.commit()

            return {
                "id": cur.lastrowid,
                "name": name,
                "gstin": gstin,
                "organization_id": current_user["organization_id"],
            }

    except Exception as e:

        raise HTTPException(
            status_code=400,
            detail=f"Could not create business: {e}",
        )


# ============================================================
# TEAM / STAFF MANAGEMENT
# ============================================================

@app.get("/api/staff")
def list_staff(current_user=Depends(get_current_user)):
    if current_user["role"] != "OWNER":
        raise HTTPException(
            status_code=403,
            detail="Only owners can manage staff",
        )

    with connect() as c:
        rows = c.execute(
            """
            SELECT
                u.id,
                u.username,
                u.role,
                u.active,
                u.created_at,
                u.last_login,
                COUNT(uba.business_id) AS assigned_businesses
            FROM users u
            LEFT JOIN user_business_access uba
                ON uba.user_id = u.id
            WHERE u.organization_id = ?
              AND u.role = 'STAFF'
            GROUP BY
                u.id,
                u.username,
                u.role,
                u.active,
                u.created_at,
                u.last_login
            ORDER BY u.username
            """,
            (current_user["organization_id"],),
        ).fetchall()

        return [dict(row) for row in rows]


@app.post("/api/staff")
def create_staff(
    username: str = Form(...),
    password: str = Form(...),
    current_user=Depends(get_current_user),
):
    if current_user["role"] != "OWNER":
        raise HTTPException(
            status_code=403,
            detail="Only owners can create staff accounts",
        )

    username = username.strip()

    if not username:
        raise HTTPException(
            status_code=400,
            detail="Username is required",
        )

    if not password:
        raise HTTPException(
            status_code=400,
            detail="Password is required",
        )

    if len(password) < 8:
        raise HTTPException(
            status_code=400,
            detail="Password must be at least 8 characters",
        )

    password_hash = hash_password(password)

    with connect() as c:
        existing = c.execute(
            """
            SELECT id
            FROM users
            WHERE organization_id = ?
              AND username = ?
            """,
            (
                current_user["organization_id"],
                username,
            ),
        ).fetchone()

        if existing:
            raise HTTPException(
                status_code=409,
                detail="Username already exists in this organization",
            )

        cursor = c.execute(
            """
            INSERT INTO users (
                organization_id,
                username,
                password_hash,
                role,
                active
            )
            VALUES (?, ?, ?, 'STAFF', 1)
            """,
            (
                current_user["organization_id"],
                username,
                password_hash,
            ),
        )

        staff_id = cursor.lastrowid

        c.execute(
            """
            INSERT INTO activity_logs (
                organization_id,
                user_id,
                action,
                entity_type,
                entity_id,
                details
            )
            VALUES (?, ?, ?, ?, ?, ?)
            """,
            (
                current_user["organization_id"],
                current_user["user_id"],
                "STAFF_CREATED",
                "user",
                staff_id,
                f"Created staff account: {username}",
            ),
        )

        c.commit()

        return {
            "ok": True,
            "user_id": staff_id,
            "username": username,
            "role": "STAFF",
            "organization_id": current_user["organization_id"],
            "active": True,
        }


def _owner_staff_in_same_org(conn, owner, staff_id):
    staff = conn.execute(
        """
        SELECT
            id,
            organization_id,
            username,
            role,
            active,
            created_at,
            last_login
        FROM users
        WHERE id = ?
          AND organization_id = ?
          AND role = 'STAFF'
        """,
        (
            staff_id,
            owner["organization_id"],
        ),
    ).fetchone()

    if not staff:
        raise HTTPException(
            status_code=404,
            detail="Staff member not found",
        )

    return staff


def _owner_business_in_same_org(conn, owner, business_id):
    business = conn.execute(
        """
        SELECT
            id,
            organization_id,
            name,
            gstin
        FROM businesses
        WHERE id = ?
          AND organization_id = ?
        """,
        (
            business_id,
            owner["organization_id"],
        ),
    ).fetchone()

    if not business:
        raise HTTPException(
            status_code=404,
            detail="Business not found",
        )

    return business


@app.get("/api/staff/{staff_id}")
def get_staff(
    staff_id: int,
    current_user=Depends(get_current_user),
):
    if current_user["role"] != "OWNER":
        raise HTTPException(
            status_code=403,
            detail="Only owners can view staff details",
        )

    with connect() as c:
        staff = _owner_staff_in_same_org(
            c,
            current_user,
            staff_id,
        )

        businesses = c.execute(
            """
            SELECT
                b.id,
                b.name,
                b.gstin
            FROM businesses b
            JOIN user_business_access uba
                ON uba.business_id = b.id
            WHERE uba.user_id = ?
              AND b.organization_id = ?
            ORDER BY b.name
            """,
            (
                staff_id,
                current_user["organization_id"],
            ),
        ).fetchall()

        result = dict(staff)
        result["businesses"] = [dict(row) for row in businesses]

        return result


@app.get("/api/staff/{staff_id}/businesses")
def get_staff_businesses(
    staff_id: int,
    current_user=Depends(get_current_user),
):
    if current_user["role"] != "OWNER":
        raise HTTPException(
            status_code=403,
            detail="Only owners can view staff assignments",
        )

    with connect() as c:
        _owner_staff_in_same_org(
            c,
            current_user,
            staff_id,
        )

        rows = c.execute(
            """
            SELECT
                b.id,
                b.name,
                b.gstin
            FROM businesses b
            JOIN user_business_access uba
                ON uba.business_id = b.id
            WHERE uba.user_id = ?
              AND b.organization_id = ?
            ORDER BY b.name
            """,
            (
                staff_id,
                current_user["organization_id"],
            ),
        ).fetchall()

        return [dict(row) for row in rows]


@app.post("/api/staff/{staff_id}/businesses/{business_id}")
def assign_staff_business(
    staff_id: int,
    business_id: int,
    current_user=Depends(get_current_user),
):
    if current_user["role"] != "OWNER":
        raise HTTPException(
            status_code=403,
            detail="Only owners can assign businesses",
        )

    with connect() as c:
        staff = _owner_staff_in_same_org(
            c,
            current_user,
            staff_id,
        )

        business = _owner_business_in_same_org(
            c,
            current_user,
            business_id,
        )

        existing = c.execute(
            """
            SELECT 1
            FROM user_business_access
            WHERE user_id = ?
              AND business_id = ?
            """,
            (
                staff_id,
                business_id,
            ),
        ).fetchone()

        if existing:
            return {
                "ok": True,
                "already_assigned": True,
                "staff_id": staff["id"],
                "business_id": business["id"],
            }

        c.execute(
            """
            INSERT INTO user_business_access (
                user_id,
                business_id
            )
            VALUES (?, ?)
            """,
            (
                staff_id,
                business_id,
            ),
        )

        c.execute(
            """
            INSERT INTO activity_logs (
                organization_id,
                user_id,
                action,
                entity_type,
                entity_id,
                details
            )
            VALUES (?, ?, ?, ?, ?, ?)
            """,
            (
                current_user["organization_id"],
                current_user["user_id"],
                "STAFF_BUSINESS_ASSIGNED",
                "business",
                business_id,
                f"Assigned business '{business['name']}' to staff '{staff['username']}'",
            ),
        )

        c.commit()

        return {
            "ok": True,
            "staff_id": staff_id,
            "business_id": business_id,
            "assigned": True,
        }


@app.delete("/api/staff/{staff_id}/businesses/{business_id}")
def unassign_staff_business(
    staff_id: int,
    business_id: int,
    current_user=Depends(get_current_user),
):
    if current_user["role"] != "OWNER":
        raise HTTPException(
            status_code=403,
            detail="Only owners can remove business assignments",
        )

    with connect() as c:
        staff = _owner_staff_in_same_org(
            c,
            current_user,
            staff_id,
        )

        business = _owner_business_in_same_org(
            c,
            current_user,
            business_id,
        )

        existing = c.execute(
            """
            SELECT 1
            FROM user_business_access
            WHERE user_id = ?
              AND business_id = ?
            """,
            (
                staff_id,
                business_id,
            ),
        ).fetchone()

        if not existing:
            return {
                "ok": True,
                "already_unassigned": True,
                "staff_id": staff_id,
                "business_id": business_id,
            }

        c.execute(
            """
            DELETE FROM user_business_access
            WHERE user_id = ?
              AND business_id = ?
            """,
            (
                staff_id,
                business_id,
            ),
        )

        c.execute(
            """
            INSERT INTO activity_logs (
                organization_id,
                user_id,
                action,
                entity_type,
                entity_id,
                details
            )
            VALUES (?, ?, ?, ?, ?, ?)
            """,
            (
                current_user["organization_id"],
                current_user["user_id"],
                "STAFF_BUSINESS_UNASSIGNED",
                "business",
                business_id,
                f"Removed business '{business['name']}' from staff '{staff['username']}'",
            ),
        )

        c.commit()

        return {
            "ok": True,
            "staff_id": staff_id,
            "business_id": business_id,
            "assigned": False,
        }


@app.post("/api/staff/{staff_id}/deactivate")
def deactivate_staff(
    staff_id: int,
    current_user=Depends(get_current_user),
):
    if current_user["role"] != "OWNER":
        raise HTTPException(
            status_code=403,
            detail="Only owners can deactivate staff",
        )

    with connect() as c:
        staff = _owner_staff_in_same_org(
            c,
            current_user,
            staff_id,
        )

        c.execute(
            """
            UPDATE users
            SET active = 0
            WHERE id = ?
              AND organization_id = ?
              AND role = 'STAFF'
            """,
            (
                staff_id,
                current_user["organization_id"],
            ),
        )

        c.execute(
            """
            DELETE FROM sessions
            WHERE user_id = ?
            """,
            (staff_id,),
        )

        c.execute(
            """
            INSERT INTO activity_logs (
                organization_id,
                user_id,
                action,
                entity_type,
                entity_id,
                details
            )
            VALUES (?, ?, ?, ?, ?, ?)
            """,
            (
                current_user["organization_id"],
                current_user["user_id"],
                "STAFF_DEACTIVATED",
                "user",
                staff_id,
                f"Deactivated staff account: {staff['username']}",
            ),
        )

        c.commit()

        return {
            "ok": True,
            "user_id": staff_id,
            "username": staff["username"],
            "active": False,
        }


# ============================================================
# CREATE ACCOUNT
# ============================================================

@app.post("/api/register")
def register(
    response: Response,
    organization_name: str = Form(...),
    username: str = Form(...),
    password: str = Form(...),
    confirm_password: str = Form(...),
):
    organization_name = organization_name.strip()
    username = username.strip()

    if not organization_name:
        raise HTTPException(
            status_code=400,
            detail="Organization name is required",
        )

    if not username:
        raise HTTPException(
            status_code=400,
            detail="Username is required",
        )

    if not password:
        raise HTTPException(
            status_code=400,
            detail="Password is required",
        )

    if password != confirm_password:
        raise HTTPException(
            status_code=400,
            detail="Passwords do not match",
        )

    if len(password) < 8:
        raise HTTPException(
            status_code=400,
            detail="Password must be at least 8 characters",
        )

    password_hash = hash_password(password)

    with connect() as c:
        try:
            cursor = c.execute(
                """
                INSERT INTO organizations (name)
                VALUES (?)
                """,
                (organization_name,),
            )

            organization_id = cursor.lastrowid

            cursor = c.execute(
                """
                INSERT INTO users (
                    organization_id,
                    username,
                    password_hash,
                    role,
                    active
                )
                VALUES (?, ?, ?, 'OWNER', 1)
                """,
                (
                    organization_id,
                    username,
                    password_hash,
                ),
            )

            user_id = cursor.lastrowid

            session_token = create_session(
                c,
                user_id,
            )

            c.commit()

        except Exception:
            c.rollback()
            raise

    response.set_cookie(
        key="session",
        value=session_token,
        httponly=True,
        samesite="lax",
        secure=False,
        max_age=12 * 60 * 60,
    )

    return {
        "ok": True,
        "username": username,
        "role": "OWNER",
        "organization_id": organization_id,
    }


# ============================================================
# INVOICE UPLOAD + PROCESSING
# ============================================================

@app.post("/api/invoices/upload")
async def upload_invoice(
    file: UploadFile = File(...),
    current_user=Depends(get_current_user),
):
    # --------------------------------------------------------
    # Validate file type
    # --------------------------------------------------------

    ext = Path(
        file.filename or "upload"
    ).suffix.lower()

    allowed_extensions = {
        ".png",
        ".jpg",
        ".jpeg",
        ".webp",
        ".pdf",
    }

    if ext not in allowed_extensions:
        raise HTTPException(
            400,
            "Supported files: PNG, JPG, WEBP, PDF"
        )

    # --------------------------------------------------------
    # Create safe filename
    # --------------------------------------------------------

    safe_name = Path(
        file.filename or "upload"
    ).name.replace(" ", "_")

    target = UPLOADS / safe_name

    counter = 1

    while target.exists():
        target = UPLOADS / (
            f"{Path(safe_name).stem}_{counter}{ext}"
        )
        counter += 1

    # --------------------------------------------------------
    # Save uploaded file
    # --------------------------------------------------------

    with target.open("wb") as out:
        shutil.copyfileobj(
            file.file,
            out
        )

    # --------------------------------------------------------
    # OCR + Parser
    # --------------------------------------------------------

    try:
        text = extract_text(target)
        data = parse_invoice(text)

        # ----------------------------------------------------
        # Deterministic verification
        # ----------------------------------------------------

        verification = verify_invoice(data)

        data["status"] = verification.status
        data["review_reasons"] = verification.reasons
        data["verification_checks"] = verification.checks

    except Exception as e:
        raise HTTPException(
            500,
            f"Processing failed: {e}"
        )

    # --------------------------------------------------------
    # MATCH INVOICE TO AUTHORIZED BUSINESS
    # --------------------------------------------------------

    business_id = None
    invoice_type = None

    with connect() as c:

        if current_user["role"] == "OWNER":

            registered_businesses = c.execute(
                """
                SELECT id, name, gstin
                FROM businesses
                WHERE organization_id = ?
                ORDER BY id
                """,
                (
                    current_user["organization_id"],
                )
            ).fetchall()

        else:

            registered_businesses = c.execute(
                """
                SELECT
                    b.id,
                    b.name,
                    b.gstin
                FROM businesses b
                JOIN user_business_access uba
                    ON uba.business_id = b.id
                WHERE b.organization_id = ?
                  AND uba.user_id = ?
                ORDER BY b.id
                """,
                (
                    current_user["organization_id"],
                    current_user["user_id"],
                )
            ).fetchall()

        invoice_text_upper = text.upper()

        for business in registered_businesses:

            business_gstin = (
                business["gstin"] or ""
            ).strip().upper()

            if not business_gstin:
                continue

            if business_gstin in invoice_text_upper:

                business_id = business["id"]

                invoice_gstins = [
                    gstin.upper()
                    for gstin in GSTIN_RE.findall(text)
                ]

                invoice_gstins = list(
                    dict.fromkeys(invoice_gstins)
                )

                if len(invoice_gstins) >= 2:

                    supplier_gstin = invoice_gstins[0]
                    buyer_gstin = invoice_gstins[1]

                    if business_gstin == supplier_gstin:
                        invoice_type = "SALES"

                    elif business_gstin == buyer_gstin:
                        invoice_type = "PURCHASE"

                    else:
                        invoice_type = None

                else:

                    invoice_type = None

                data["gstin"] = business_gstin

                break

        # ----------------------------------------------------
        # Do not save an invoice if no authorized business
        # matches the invoice.
        # ----------------------------------------------------

        if business_id is None:

            raise HTTPException(
                400,
                "Invoice could not be matched to an authorized business"
            )

        # ----------------------------------------------------
        # Save invoice + line items
        # ----------------------------------------------------

        cur = c.execute(
            """
            INSERT INTO invoices(
                business_id,
                organization_id,
                source_filename,
                supplier_name,
                gstin,
                invoice_type,
                invoice_number,
                invoice_date,
                taxable_value,
                cgst,
                sgst,
                igst,
                total_amount,
                confidence,
                status,
                review_reasons,
                verification_checks,
                raw_text,
                uploaded_by
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                business_id,
                current_user["organization_id"],
                target.name,
                data.get("supplier_name"),
                data.get("gstin"),
                invoice_type,
                data.get("invoice_number"),
                data.get("invoice_date"),
                data.get("taxable_value"),
                data.get("cgst"),
                data.get("sgst"),
                data.get("igst"),
                data.get("total_amount"),
                data.get("confidence"),
                data.get("status"),
                json.dumps(
                    data.get("review_reasons", []),
                    ensure_ascii=False,
                ),
                json.dumps(
                    data.get("verification_checks", {}),
                    ensure_ascii=False,
                ),
                text,
                current_user["user_id"],
            )
        )

        invoice_id = cur.lastrowid

        # ----------------------------------------------------
        # Save invoice line items
        # ----------------------------------------------------

        for item in data.get("items", []):

            c.execute(
                """
                INSERT INTO invoice_items(
                    invoice_id,
                    serial_no,
                    description,
                    hsn_sac,
                    quantity,
                    rate,
                    taxable_value
                )
                VALUES (?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    invoice_id,
                    item.get("serial_no"),
                    item.get("description"),
                    item.get("hsn_sac"),
                    item.get("quantity"),
                    item.get("rate"),
                    item.get("taxable_value"),
                )
            )

        c.commit()

    # --------------------------------------------------------
    # Return result to UI
    # --------------------------------------------------------

    return {
        "id": invoice_id,
        "business_id": business_id,
        **data,
    }


# ============================================================
# INVOICES / REVIEW
# ============================================================

class ReviewActionRequest(BaseModel):
    notes: str | None = None


def _review_authorized_invoice(conn, invoice_id: int, current_user):
    if current_user["role"] == "OWNER":
        row = conn.execute(
            """
            SELECT
                i.id,
                i.organization_id,
                i.status,
                i.review_status,
                b.id AS business_id,
                b.name AS business_name,
                b.organization_id AS business_organization_id
            FROM invoices i
            JOIN businesses b
                ON b.id = i.business_id
            WHERE i.id = ?
              AND i.organization_id = ?
              AND b.organization_id = ?
            """,
            (
                invoice_id,
                current_user["organization_id"],
                current_user["organization_id"],
            ),
        ).fetchone()
    else:
        row = conn.execute(
            """
            SELECT
                i.id,
                i.organization_id,
                i.status,
                i.review_status,
                b.id AS business_id,
                b.name AS business_name,
                b.organization_id AS business_organization_id
            FROM invoices i
            JOIN businesses b
                ON b.id = i.business_id
            JOIN user_business_access uba
                ON uba.business_id = b.id
            WHERE i.id = ?
              AND i.organization_id = ?
              AND b.organization_id = ?
              AND uba.user_id = ?
            """,
            (
                invoice_id,
                current_user["organization_id"],
                current_user["organization_id"],
                current_user["user_id"],
            ),
        ).fetchone()

    if not row:
        raise HTTPException(
            status_code=404,
            detail="Invoice not found",
        )

    return row


@app.post("/api/invoices/{invoice_id}/review/approve")
def approve_invoice_review(
    invoice_id: int,
    payload: ReviewActionRequest,
    current_user=Depends(get_current_user),
):
    with connect() as conn:
        invoice = _review_authorized_invoice(
            conn,
            invoice_id,
            current_user,
        )

        if invoice["review_status"] != "pending":
            raise HTTPException(
                status_code=409,
                detail="Invoice has already been reviewed",
            )

        reviewed_at = datetime.now(timezone.utc).isoformat()

        conn.execute(
            """
            UPDATE invoices
            SET
                review_status = 'approved',
                reviewed_by = ?,
                reviewed_at = ?,
                review_notes = ?
            WHERE id = ?
              AND organization_id = ?
            """,
            (
                current_user["user_id"],
                reviewed_at,
                payload.notes,
                invoice_id,
                current_user["organization_id"],
            ),
        )

        conn.commit()

        return {
            "id": invoice_id,
            "business_id": invoice["business_id"],
            "business_name": invoice["business_name"],
            "status": invoice["status"],
            "review_status": "approved",
            "reviewed_by": current_user["user_id"],
            "reviewed_at": reviewed_at,
            "review_notes": payload.notes,
        }


@app.post("/api/invoices/{invoice_id}/review/reject")
def reject_invoice_review(
    invoice_id: int,
    payload: ReviewActionRequest,
    current_user=Depends(get_current_user),
):
    with connect() as conn:
        invoice = _review_authorized_invoice(
            conn,
            invoice_id,
            current_user,
        )

        if invoice["review_status"] != "pending":
            raise HTTPException(
                status_code=409,
                detail="Invoice has already been reviewed",
            )

        reviewed_at = datetime.now(timezone.utc).isoformat()

        conn.execute(
            """
            UPDATE invoices
            SET
                review_status = 'rejected',
                reviewed_by = ?,
                reviewed_at = ?,
                review_notes = ?
            WHERE id = ?
              AND organization_id = ?
            """,
            (
                current_user["user_id"],
                reviewed_at,
                payload.notes,
                invoice_id,
                current_user["organization_id"],
            ),
        )

        conn.commit()

        return {
            "id": invoice_id,
            "business_id": invoice["business_id"],
            "business_name": invoice["business_name"],
            "status": invoice["status"],
            "review_status": "rejected",
            "reviewed_by": current_user["user_id"],
            "reviewed_at": reviewed_at,
            "review_notes": payload.notes,
        }


@app.get("/api/invoices/{invoice_id}/review")
def review_invoice_detail(
    invoice_id: int,
    current_user=Depends(get_current_user),
):
    with connect() as c:

        if current_user["role"] == "OWNER":

            row = c.execute(
                """
                SELECT
                    i.id,
                    i.business_id,
                    i.organization_id,
                    b.name AS business_name,
                    b.gstin AS business_gstin,
                    i.source_filename,
                    i.supplier_name,
                    i.gstin,
                    i.invoice_number,
                    i.invoice_date,
                    i.taxable_value,
                    i.cgst,
                    i.sgst,
                    i.igst,
                    i.total_amount,
                    i.confidence,
                    i.status,
                    i.review_reasons,
                    i.verification_checks,
                    i.raw_text,
                    i.created_at
                FROM invoices i
                JOIN businesses b
                    ON b.id = i.business_id
                WHERE i.id = ?
                  AND i.organization_id = ?
                  AND b.organization_id = ?
                """,
                (
                    invoice_id,
                    current_user["organization_id"],
                    current_user["organization_id"],
                ),
            ).fetchone()

        else:

            row = c.execute(
                """
                SELECT
                    i.id,
                    i.business_id,
                    i.organization_id,
                    b.name AS business_name,
                    b.gstin AS business_gstin,
                    i.source_filename,
                    i.supplier_name,
                    i.gstin,
                    i.invoice_number,
                    i.invoice_date,
                    i.taxable_value,
                    i.cgst,
                    i.sgst,
                    i.igst,
                    i.total_amount,
                    i.confidence,
                    i.status,
                    i.review_reasons,
                    i.verification_checks,
                    i.raw_text,
                    i.created_at
                FROM invoices i
                JOIN businesses b
                    ON b.id = i.business_id
                JOIN user_business_access uba
                    ON uba.business_id = b.id
                WHERE i.id = ?
                  AND i.organization_id = ?
                  AND b.organization_id = ?
                  AND uba.user_id = ?
                """,
                (
                    invoice_id,
                    current_user["organization_id"],
                    current_user["organization_id"],
                    current_user["user_id"],
                ),
            ).fetchone()

        if row is None:
            raise HTTPException(
                status_code=404,
                detail="Invoice not found",
            )

        invoice = dict(row)

        try:
            invoice["review_reasons"] = json.loads(
                invoice["review_reasons"] or "[]"
            )
        except (TypeError, json.JSONDecodeError):
            invoice["review_reasons"] = []

        try:
            invoice["verification_checks"] = json.loads(
                invoice["verification_checks"] or "{}"
            )
        except (TypeError, json.JSONDecodeError):
            invoice["verification_checks"] = {}

        items = c.execute(
            """
            SELECT
                id,
                serial_no,
                description,
                hsn_sac,
                quantity,
                rate,
                taxable_value,
                created_at
            FROM invoice_items
            WHERE invoice_id = ?
            ORDER BY
                serial_no ASC,
                id ASC
            """,
            (invoice_id,),
        ).fetchall()

        invoice["items"] = [
            dict(item)
            for item in items
        ]

        return invoice


@app.get("/api/invoices/review")
def review_queue(
    current_user=Depends(get_current_user),
):
    with connect() as c:

        if current_user["role"] == "OWNER":

            rows = c.execute(
                """
                SELECT
                    i.id,
                    i.business_id,
                    b.name AS business_name,
                    b.gstin AS business_gstin,
                    i.source_filename,
                    i.supplier_name,
                    i.gstin,
                    i.invoice_number,
                    i.invoice_date,
                    i.taxable_value,
                    i.cgst,
                    i.sgst,
                    i.igst,
                    i.total_amount,
                    i.confidence,
                    i.status,
                    i.review_reasons,
                    i.verification_checks,
                    i.created_at
                FROM invoices i
                JOIN businesses b
                    ON b.id = i.business_id
                WHERE i.organization_id = ?
                  AND i.status = 'needs_review'
                  AND b.organization_id = ?
                ORDER BY i.created_at DESC, i.id DESC
                """,
                (
                    current_user["organization_id"],
                    current_user["organization_id"],
                ),
            ).fetchall()

        else:

            rows = c.execute(
                """
                SELECT
                    i.id,
                    i.business_id,
                    b.name AS business_name,
                    b.gstin AS business_gstin,
                    i.source_filename,
                    i.supplier_name,
                    i.gstin,
                    i.invoice_number,
                    i.invoice_date,
                    i.taxable_value,
                    i.cgst,
                    i.sgst,
                    i.igst,
                    i.total_amount,
                    i.confidence,
                    i.status,
                    i.review_reasons,
                    i.verification_checks,
                    i.created_at
                FROM invoices i
                JOIN businesses b
                    ON b.id = i.business_id
                JOIN user_business_access uba
                    ON uba.business_id = b.id
                WHERE i.organization_id = ?
                  AND b.organization_id = ?
                  AND uba.user_id = ?
                  AND i.status = 'needs_review'
                ORDER BY i.created_at DESC, i.id DESC
                """,
                (
                    current_user["organization_id"],
                    current_user["organization_id"],
                    current_user["user_id"],
                ),
            ).fetchall()

        results = []

        for row in rows:

            item = dict(row)

            try:
                item["review_reasons"] = json.loads(
                    item["review_reasons"] or "[]"
                )
            except (TypeError, json.JSONDecodeError):
                item["review_reasons"] = []

            try:
                item["verification_checks"] = json.loads(
                    item["verification_checks"] or "{}"
                )
            except (TypeError, json.JSONDecodeError):
                item["verification_checks"] = {}

            results.append(item)

        return results


@app.get("/api/invoices")
def invoices(
    business_id: int | None = None,
    current_user=Depends(get_current_user),
):

    with connect() as c:

        if current_user["role"] == "OWNER":

            if business_id is not None:

                rows = c.execute(
                    """
                    SELECT
                        i.*,
                        b.name AS business_name,
                        b.gstin AS business_gstin
                    FROM invoices i
                    JOIN businesses b
                        ON b.id = i.business_id
                    WHERE i.business_id = ?
                      AND b.organization_id = ?
                    ORDER BY i.id DESC
                    """,
                    (
                        business_id,
                        current_user["organization_id"],
                    )
                ).fetchall()

            else:

                rows = c.execute(
                    """
                    SELECT
                        i.*,
                        b.name AS business_name,
                        b.gstin AS business_gstin
                    FROM invoices i
                    JOIN businesses b
                        ON b.id = i.business_id
                    WHERE b.organization_id = ?
                    ORDER BY i.id DESC
                    """,
                    (
                        current_user["organization_id"],
                    )
                ).fetchall()

        else:

            if business_id is not None:

                rows = c.execute(
                    """
                    SELECT
                        i.*,
                        b.name AS business_name,
                        b.gstin AS business_gstin
                    FROM invoices i
                    JOIN businesses b
                        ON b.id = i.business_id
                    JOIN user_business_access uba
                        ON uba.business_id = b.id
                    WHERE i.business_id = ?
                      AND b.organization_id = ?
                      AND uba.user_id = ?
                    ORDER BY i.id DESC
                    """,
                    (
                        business_id,
                        current_user["organization_id"],
                        current_user["user_id"],
                    )
                ).fetchall()

            else:

                rows = c.execute(
                    """
                    SELECT
                        i.*,
                        b.name AS business_name,
                        b.gstin AS business_gstin
                    FROM invoices i
                    JOIN businesses b
                        ON b.id = i.business_id
                    JOIN user_business_access uba
                        ON uba.business_id = b.id
                    WHERE b.organization_id = ?
                      AND uba.user_id = ?
                    ORDER BY i.id DESC
                    """,
                    (
                        current_user["organization_id"],
                        current_user["user_id"],
                    )
                ).fetchall()

        return [
            dict(row)
            for row in rows
        ]


# ============================================================
# INVOICE ITEMS
# ============================================================

@app.get("/api/invoices/{invoice_id}/items")
def invoice_items(
    invoice_id: int,
    current_user=Depends(get_current_user),
):

    with connect() as c:

        if current_user["role"] == "OWNER":

            invoice = c.execute(
                """
                SELECT
                    i.id,
                    i.business_id,
                    b.organization_id
                FROM invoices i
                JOIN businesses b
                    ON b.id = i.business_id
                WHERE i.id = ?
                  AND b.organization_id = ?
                """,
                (
                    invoice_id,
                    current_user["organization_id"],
                )
            ).fetchone()

        else:

            invoice = c.execute(
                """
                SELECT
                    i.id,
                    i.business_id,
                    b.organization_id
                FROM invoices i
                JOIN businesses b
                    ON b.id = i.business_id
                JOIN user_business_access uba
                    ON uba.business_id = b.id
                WHERE i.id = ?
                  AND b.organization_id = ?
                  AND uba.user_id = ?
                """,
                (
                    invoice_id,
                    current_user["organization_id"],
                    current_user["user_id"],
                )
            ).fetchone()

        if not invoice:
            raise HTTPException(
                status_code=404,
                detail="Invoice not found",
            )

        rows = c.execute(
            """
            SELECT
                id,
                invoice_id,
                serial_no,
                description,
                hsn_sac,
                quantity,
                rate,
                taxable_value
            FROM invoice_items
            WHERE invoice_id = ?
            ORDER BY serial_no
            """,
            (invoice_id,)
        ).fetchall()

        return [
            dict(row)
            for row in rows
        ]


# ============================================================
# APPROVE INVOICE
# ============================================================

@app.post("/api/invoices/{invoice_id}/approve")
def approve(
    invoice_id: int,
    current_user=Depends(get_current_user),
):

    with connect() as c:

        if current_user["role"] == "OWNER":

            invoice = c.execute(
                """
                SELECT
                    i.id,
                    i.business_id
                FROM invoices i
                JOIN businesses b
                    ON b.id = i.business_id
                WHERE i.id = ?
                  AND b.organization_id = ?
                """,
                (
                    invoice_id,
                    current_user["organization_id"],
                )
            ).fetchone()

        else:

            invoice = c.execute(
                """
                SELECT
                    i.id,
                    i.business_id
                FROM invoices i
                JOIN businesses b
                    ON b.id = i.business_id
                JOIN user_business_access uba
                    ON uba.business_id = b.id
                WHERE i.id = ?
                  AND b.organization_id = ?
                  AND uba.user_id = ?
                """,
                (
                    invoice_id,
                    current_user["organization_id"],
                    current_user["user_id"],
                )
            ).fetchone()

        if not invoice:
            raise HTTPException(
                status_code=404,
                detail="Invoice not found",
            )

        # ----------------------------------------------------
        # Approve invoice
        # ----------------------------------------------------

        cur = c.execute(
            """
            UPDATE invoices
            SET status = 'approved'
            WHERE id = ?
              AND organization_id = ?
            """,
            (
                invoice_id,
                current_user["organization_id"],
            )
        )

        c.commit()

        if not cur.rowcount:
            raise HTTPException(
                status_code=404,
                detail="Invoice not found",
            )

    return {
        "ok": True
    }


# ============================================================
# EXCEL EXPORT
# ============================================================

@app.get("/api/export.xlsx")
def export_excel(
    business_id: int | None = None,
    current_user=Depends(get_current_user),
):
    with connect() as c:

        if current_user["role"] == "OWNER":

            if business_id is not None:

                businesses = c.execute(
                    """
                    SELECT
                        id,
                        name,
                        gstin
                    FROM businesses
                    WHERE id = ?
                      AND organization_id = ?
                    ORDER BY name
                    """,
                    (
                        business_id,
                        current_user["organization_id"],
                    )
                ).fetchall()

            else:

                businesses = c.execute(
                    """
                    SELECT
                        id,
                        name,
                        gstin
                    FROM businesses
                    WHERE organization_id = ?
                    ORDER BY name
                    """,
                    (
                        current_user["organization_id"],
                    )
                ).fetchall()

        else:

            if business_id is not None:

                businesses = c.execute(
                    """
                    SELECT
                        b.id,
                        b.name,
                        b.gstin
                    FROM businesses b
                    JOIN user_business_access uba
                        ON uba.business_id = b.id
                    WHERE b.id = ?
                      AND b.organization_id = ?
                      AND uba.user_id = ?
                    ORDER BY b.name
                    """,
                    (
                        business_id,
                        current_user["organization_id"],
                        current_user["user_id"],
                    )
                ).fetchall()

            else:

                businesses = c.execute(
                    """
                    SELECT
                        b.id,
                        b.name,
                        b.gstin
                    FROM businesses b
                    JOIN user_business_access uba
                        ON uba.business_id = b.id
                    WHERE b.organization_id = ?
                      AND uba.user_id = ?
                    ORDER BY b.name
                    """,
                    (
                        current_user["organization_id"],
                        current_user["user_id"],
                    )
                ).fetchall()

        if business_id is not None and not businesses:
            raise HTTPException(
                status_code=404,
                detail="Business not found",
            )

        wb = Workbook()

        default_sheet = wb.active
        wb.remove(default_sheet)

        for business in businesses:

            business_id_value = business["id"]

            sheet_name = (
                business["name"]
                or f"Business {business_id_value}"
            )

            for char in ['\\', '/', '*', '?', ':', '[', ']']:
                sheet_name = sheet_name.replace(char, '-')

            sheet_name = sheet_name[:31].strip()

            if not sheet_name:
                sheet_name = f"Business {business_id_value}"

            original_sheet_name = sheet_name
            counter = 2

            while sheet_name in wb.sheetnames:

                suffix = f" ({counter})"

                sheet_name = (
                    original_sheet_name[:31 - len(suffix)]
                    + suffix
                )

                counter += 1

            ws = wb.create_sheet(sheet_name)

            # ---------------------------------------------------------
            # BUSINESS HEADER
            # ---------------------------------------------------------

            ws["A1"] = "Business"
            ws["B1"] = business["name"]

            ws["A2"] = "GSTIN"
            ws["B2"] = business["gstin"]

            ws["A4"] = "Invoice Summary"

            for cell in ["A1", "A2", "A4"]:
                ws[cell].font = Font(bold=True)

            # ---------------------------------------------------------
            # INVOICE SUMMARY
            # ---------------------------------------------------------

            invoice_headers = [
                "ID",
                "Supplier",
                "GSTIN",
                "Invoice No",
                "Date",
                "Taxable Value",
                "CGST",
                "SGST",
                "IGST",
                "Total",
                "Confidence",
                "Status",
            ]

            header_row = 5

            for col, header in enumerate(
                invoice_headers,
                start=1
            ):

                cell = ws.cell(
                    row=header_row,
                    column=col,
                    value=header
                )

                cell.font = Font(bold=True)

            # ---------------------------------------------------------
            # INVOICES
            # ---------------------------------------------------------

            if current_user["role"] == "OWNER":

                invoices = c.execute(
                    """
                    SELECT
                        i.*
                    FROM invoices i
                    JOIN businesses b
                        ON b.id = i.business_id
                    WHERE i.business_id = ?
                      AND b.organization_id = ?
                    ORDER BY i.id
                    """,
                    (
                        business_id_value,
                        current_user["organization_id"],
                    )
                ).fetchall()

            else:

                invoices = c.execute(
                    """
                    SELECT
                        i.*
                    FROM invoices i
                    JOIN businesses b
                        ON b.id = i.business_id
                    JOIN user_business_access uba
                        ON uba.business_id = b.id
                    WHERE i.business_id = ?
                      AND b.organization_id = ?
                      AND uba.user_id = ?
                    ORDER BY i.id
                    """,
                    (
                        business_id_value,
                        current_user["organization_id"],
                        current_user["user_id"],
                    )
                ).fetchall()

            row_number = header_row + 1

            invoice_ids = []

            for invoice in invoices:

                invoice_ids.append(invoice["id"])

                ws.append([
                    invoice["id"],
                    invoice["supplier_name"],
                    invoice["gstin"],
                    invoice["invoice_number"],
                    invoice["invoice_date"],
                    invoice["taxable_value"],
                    invoice["cgst"],
                    invoice["sgst"],
                    invoice["igst"],
                    invoice["total_amount"],
                    invoice["confidence"],
                    invoice["status"],
                ])

                row_number += 1

            # ---------------------------------------------------------
            # LINE ITEMS
            # ---------------------------------------------------------

            items_start_row = row_number + 2

            ws.cell(
                row=items_start_row,
                column=1,
                value="Invoice Items"
            ).font = Font(bold=True)

            item_header_row = items_start_row + 1

            item_headers = [
                "Invoice ID",
                "Invoice No",
                "Serial No",
                "Description",
                "HSN/SAC",
                "Quantity",
                "Rate",
                "Taxable Value",
            ]

            for col, header in enumerate(
                item_headers,
                start=1
            ):

                cell = ws.cell(
                    row=item_header_row,
                    column=col,
                    value=header
                )

                cell.font = Font(bold=True)

            if invoice_ids:

                placeholders = ",".join(
                    "?" for _ in invoice_ids
                )

                items = c.execute(
                    f"""
                    SELECT
                        ii.invoice_id,
                        i.invoice_number,
                        ii.serial_no,
                        ii.description,
                        ii.hsn_sac,
                        ii.quantity,
                        ii.rate,
                        ii.taxable_value
                    FROM invoice_items ii
                    JOIN invoices i
                        ON i.id = ii.invoice_id
                    WHERE ii.invoice_id IN ({placeholders})
                    ORDER BY ii.invoice_id, ii.serial_no
                    """,
                    invoice_ids
                ).fetchall()

                for item in items:

                    ws.append([
                        item["invoice_id"],
                        item["invoice_number"],
                        item["serial_no"],
                        item["description"],
                        item["hsn_sac"],
                        item["quantity"],
                        item["rate"],
                        item["taxable_value"],
                    ])

            # ---------------------------------------------------------
            # FORMATTING
            # ---------------------------------------------------------

            ws.freeze_panes = "A6"

            for row in ws.iter_rows():

                for cell in row:

                    if cell.column in [
                        6, 7, 8, 9, 10
                    ] and isinstance(
                        cell.value,
                        (int, float)
                    ):

                        cell.number_format = '?#,##0.00'

            for row in range(
                item_header_row + 1,
                ws.max_row + 1
            ):

                ws.cell(
                    row=row,
                    column=7
                ).number_format = '?#,##0.00'

                ws.cell(
                    row=row,
                    column=8
                ).number_format = '?#,##0.00'

            for row in range(
                header_row + 1,
                items_start_row
            ):

                ws.cell(
                    row=row,
                    column=11
                ).number_format = '0%'

            widths = {
                "A": 12,
                "B": 30,
                "C": 20,
                "D": 20,
                "E": 14,
                "F": 16,
                "G": 14,
                "H": 14,
                "I": 14,
                "J": 16,
                "K": 14,
                "L": 14,
            }

            for column, width in widths.items():
                ws.column_dimensions[column].width = width

        if not businesses:

            ws = wb.create_sheet("No Businesses")

            ws["A1"] = (
                "No authorized client businesses available."
            )

            ws["A1"].font = Font(bold=True)

            ws.column_dimensions["A"].width = 50

    output_file = EXPORTS / "gstflow_invoices.xlsx"

    wb.save(output_file)

    return FileResponse(
        output_file,
        filename="gstflow_invoices.xlsx"
    )